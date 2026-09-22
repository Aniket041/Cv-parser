"""Local PDF/DOCX text extraction and OCR fallback."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass

# Wingdings/Symbol bullets survive PDF extraction as Private Use Area code points
# and end up glued to the value after them (" Deputy General Manager").
JUNK_CHARS = re.compile(r"[-�\x00-\x08\x0b\x0c\x0e-\x1f]")


def clean_text(text: str) -> str:
    return JUNK_CHARS.sub(" ", text or "")


@dataclass
class ExtractedText:
    text: str
    pages: list[str]
    used_ocr: bool = False
    warnings: list[str] | None = None


def extract_docx(content: bytes) -> ExtractedText:
    from docx import Document

    document = Document(io.BytesIO(content))
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table_text: list[str] = []
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            if any(cells):
                table_text.append(" | ".join(cells))
    text = clean_text("\n".join(paragraphs + table_text))
    return ExtractedText(text=text, pages=[text], warnings=[])


def extract_pdf(content: bytes, enable_ocr: bool = True) -> ExtractedText:
    import fitz

    pdf = fitz.open(stream=content, filetype="pdf")
    try:
        pages = [clean_text(page.get_text("text")).strip() for page in pdf]
        text = "\n\n".join(pages).strip()
        # A small amount of text may be a header/footer; use OCR for genuinely image-only resumes.
        if len(text) >= 120 or not enable_ocr:
            return ExtractedText(text=text, pages=pages, warnings=[])

        try:
            import pytesseract
            from PIL import Image
        except ImportError:
            return ExtractedText(text=text, pages=pages, warnings=["OCR Python packages are not installed."])

        ocr_pages: list[str] = []
        for page in pdf:
            pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            image = Image.open(io.BytesIO(pix.tobytes("png")))
            ocr_pages.append(clean_text(pytesseract.image_to_string(image)).strip())
        ocr_text = "\n\n".join(ocr_pages).strip()
        if ocr_text:
            return ExtractedText(text=ocr_text, pages=ocr_pages, used_ocr=True, warnings=[])
        return ExtractedText(text=text, pages=pages, warnings=["OCR returned no readable text."])
    finally:
        pdf.close()


def extract_file(filename: str, content: bytes, enable_ocr: bool = True) -> ExtractedText:
    extension = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension == "pdf":
        return extract_pdf(content, enable_ocr)
    if extension == "docx":
        return extract_docx(content)
    raise ValueError("Only PDF and DOCX files are supported.")

